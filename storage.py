"""Excel storage layer: every order is saved to data/sales_data.xlsx."""
import os
import threading
from datetime import datetime, date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
XLSX_PATH = os.path.join(DATA_DIR, "sales_data.xlsx")

HEADERS = [
    "OrderID", "Date", "CustomerID", "CustomerName", "Email", "City",
    "Product", "Category", "Qty", "UnitPrice", "Total", "Source",
]

_lock = threading.Lock()


def _style_sheet(ws):
    fill = PatternFill("solid", fgColor="1F2A44")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    widths = [18, 12, 26, 22, 30, 16, 28, 16, 6, 11, 11, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def _row_to_values(o):
    d = o["Date"]
    if isinstance(d, str):
        d = datetime.strptime(d, "%Y-%m-%d").date()
    return [
        o["OrderID"], d, o["CustomerID"], o["CustomerName"], o.get("Email", ""),
        o.get("City", ""), o["Product"], o.get("Category", "Other"),
        int(o["Qty"]), float(o["UnitPrice"]),
        round(int(o["Qty"]) * float(o["UnitPrice"]), 2), o.get("Source", "manual"),
    ]


def init(seed_orders=None):
    """Create the workbook if it does not exist (optionally seeded)."""
    os.makedirs(DATA_DIR, exist_ok=True)
    with _lock:
        if os.path.exists(XLSX_PATH):
            return False
        wb = Workbook()
        ws = wb.active
        ws.title = "Orders"
        ws.append(HEADERS)
        for o in (seed_orders or []):
            ws.append(_row_to_values(o))
        ws["A1"].value = "OrderID"
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
            row[0].number_format = "yyyy-mm-dd"
        _style_sheet(ws)
        wb.save(XLSX_PATH)
        return True


def read_orders():
    with _lock:
        wb = load_workbook(XLSX_PATH, read_only=True, data_only=True)
        ws = wb["Orders"]
        out = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not row or row[0] is None:
                continue
            rec = dict(zip(HEADERS, row))
            d = rec["Date"]
            if isinstance(d, datetime):
                d = d.date()
            elif isinstance(d, str):
                d = datetime.strptime(d[:10], "%Y-%m-%d").date()
            if not isinstance(d, date):
                continue
            rec["Date"] = d
            rec["Qty"] = int(rec["Qty"] or 0)
            rec["UnitPrice"] = float(rec["UnitPrice"] or 0)
            rec["Total"] = float(rec["Total"] or rec["Qty"] * rec["UnitPrice"])
            out.append(rec)
        wb.close()
        return out


def append_orders(new_orders):
    """Append orders, skipping OrderIDs that already exist. Returns count added."""
    if not new_orders:
        return 0
    with _lock:
        try:
            wb = load_workbook(XLSX_PATH)
        except PermissionError:
            raise PermissionError(
                "Close sales_data.xlsx in Excel first - it is locked by another program."
            )
        ws = wb["Orders"]
        existing = {r[0] for r in ws.iter_rows(min_row=2, max_col=1, values_only=True) if r[0]}
        added = 0
        for o in new_orders:
            if o["OrderID"] in existing:
                continue
            ws.append(_row_to_values(o))
            ws.cell(row=ws.max_row, column=2).number_format = "yyyy-mm-dd"
            existing.add(o["OrderID"])
            added += 1
        if added:
            try:
                wb.save(XLSX_PATH)
            except PermissionError:
                raise PermissionError(
                    "Close sales_data.xlsx in Excel first - it is locked by another program."
                )
        wb.close()
        return added
